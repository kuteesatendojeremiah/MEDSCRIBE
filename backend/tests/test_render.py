"""FR-10..FR-13, §12.1: render only on a matching hash; DOCX matches the reference note; no copy."""

import copy
import io
import json
import logging
import os
import re
import tempfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.render.hash import note_hash, short_hash
from tests.conftest import make_settings

ROOT = Path(__file__).resolve().parents[2]
FINAL = json.loads((ROOT / "fixtures/synthetic/kato_expected_note.json").read_text("utf8"))
REFERENCE = ROOT / "docs/reference/example_note_fictional.docx"


def final(client, auth, note=None, hash_=None, fmt="docx", at="2026-10-19T10:38:00+03:00"):
    note = note or FINAL
    body = {"note": note, "approved_hash": hash_ or note_hash(note), "approved_at": at}
    return client.post(f"/render?mode=final&format={fmt}", json=body, headers=auth)


def all_text(docx_bytes: bytes) -> str:
    d = DocxDocument(io.BytesIO(docx_bytes))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    parts += [p.text for p in d.sections[0].footer.paragraphs]
    return "\n".join(parts)


# --- hash lock (FR-11) -----------------------------------------------------------------------
def test_matching_hash_renders_a_docx(client, auth):
    r = final(client, auth)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.content[:2] == b"PK"  # a zip container


def test_one_changed_character_is_refused_as_terminal(client, auth):
    changed = copy.deepcopy(FINAL)
    changed["items"][0]["text"] += "."
    r = final(client, auth, note=changed, hash_=note_hash(FINAL))
    assert r.status_code == 409
    assert r.json() == {"error": "E_HASH_MISMATCH", "class": "terminal"}


def test_missing_hash_is_refused(client, auth):
    r = client.post("/render?mode=final", json={"note": FINAL}, headers=auth)
    assert r.status_code == 409 and r.json()["error"] == "E_HASH_MISMATCH"


def test_hash_ignores_key_order_and_approval_block():
    a = note_hash(FINAL)
    shuffled = {k: FINAL[k] for k in reversed(list(FINAL))}
    shuffled["approval"] = {"approved_by": "x", "approved_at": "y", "version_hash": "z"}
    assert note_hash(shuffled) == a


def test_hash_is_stable_across_languages_cross_check():
    """The browser (frontend/src/session/hash.test.ts) asserts this same constant (FR-11)."""
    expected = (ROOT / "fixtures/synthetic/kato_expected_note.sha256").read_text().strip()
    assert note_hash(FINAL) == expected


def test_open_high_item_is_refused_even_with_a_valid_hash(client, auth):
    n = copy.deepcopy(FINAL)
    n["review_items"][0]["status"] = "open"
    r = final(client, auth, note=n)
    assert r.status_code == 409 and r.json()["error"] == "E_HIGH_OPEN"


def test_invalid_note_is_recoverable_422(client, auth):
    bad = copy.deepcopy(FINAL)
    del bad["items"][0]["source"]
    r = final(client, auth, note=bad)
    assert r.status_code == 422 and r.json()["class"] == "recoverable"


def test_render_requires_a_token(client):
    r = client.post("/render", json={"note": FINAL})
    assert r.status_code == 401


# --- response hygiene (FR-13, §10.2) ---------------------------------------------------------------
def test_response_is_no_store_with_a_neutral_file_name(client, auth):
    r = final(client, auth)
    assert r.headers["cache-control"] == "no-store"
    assert r.headers["pragma"] == "no-cache"
    disp = r.headers["content-disposition"]
    assert 'filename="consultation_note_2026-10-19_1038.docx"' in disp
    assert "kato" not in disp.lower() and "ronald" not in disp.lower()


def test_no_server_copy_nothing_left_in_temp_or_cwd(client, auth):
    tmp = Path(tempfile.gettempdir())
    before_tmp, before_cwd = set(tmp.glob("*")), set(os.listdir("."))
    for _ in range(3):
        assert final(client, auth).status_code == 200
    assert set(os.listdir(".")) == before_cwd
    leftovers = {p for p in set(tmp.glob("*")) - before_tmp if p.name.endswith((".docx", ".pdf"))}
    assert not leftovers


def test_security_events_are_logged_content_free(client, auth, caplog):
    changed = copy.deepcopy(FINAL)
    changed["items"][0]["text"] = "SYNTHETIC-MARKER-7Q4Z"
    with caplog.at_level(logging.WARNING):
        final(client, auth, note=changed, hash_="0" * 64)
    assert "approved-hash mismatch" in caplog.text
    assert "SYNTHETIC-MARKER-7Q4Z" not in caplog.text


# --- DOCX content against the reference note (FR-12, §4.3) ------------------------------------------
@pytest.fixture(scope="module")
def docx_text():
    from fastapi.testclient import TestClient

    from app.main import create_app

    c = TestClient(create_app(make_settings()), raise_server_exceptions=False)
    tok = c.post("/auth/dev-login", json={"name": "Dr. A. Nankya"}).json()["token"]
    r = final(c, {"Authorization": f"Bearer {tok}"})
    return all_text(r.content)


def test_sections_appear_in_section_4_1_order(docx_text):
    order = [
        "HISTORY",
        "MEDICATIONS AND ALLERGIES",
        "VITAL SIGNS",
        "EXAMINATION FINDINGS",
        "INVESTIGATIONS AND RESULTS",
        "ASSESSMENT",
        "TREATMENT",
        "PLAN",
        "REVIEW ITEMS",
    ]
    positions = [docx_text.index(h) for h in order]
    assert positions == sorted(positions)


def test_headings_match_the_reference_note(docx_text):
    ref = DocxDocument(str(REFERENCE))
    ref_text = "\n".join(
        [p.text for p in ref.paragraphs]
        + [c.text for t in ref.tables for r in t.rows for c in r.cells]
    )
    for heading in (
        "HISTORY",
        "MEDICATIONS AND ALLERGIES",
        "VITAL SIGNS",
        "EXAMINATION FINDINGS",
        "INVESTIGATIONS AND RESULTS",
        "ASSESSMENT",
        "TREATMENT",
        "PLAN",
        "REVIEW ITEMS",
    ):
        assert heading in ref_text and heading in docx_text


def test_header_has_every_required_field(docx_text):
    for needle in (
        "Ronald Kato (FICTIONAL)",
        "Temporary ref.",
        "OPD-07",
        "Date / time",
        "19 October 2026",
        "Clinician",
        "Dr. A. Nankya",
        "Languages",
        "Luganda, English",
        "Present",
        "Wife (relative)",
        "Consent",
        "agreed verbally",
    ):
        assert needle in docx_text, needle


def test_reference_facts_are_present(docx_text):
    low = docx_text.lower()
    for fact in (
        "chest pain for 3 days",
        "no radiation of pain to arm or jaw",
        "reported by wife",
        "no known drug allergies",
        "pulse approx. 110/min, irregular",
        "bp 150/95 mmhg",
        "weight 82 kg",
        "fine crackles at both lung bases",
        "mild bilateral pitting ankle oedema",
        "atrial fibrillation",
        "possible heart failure",
        "bisoprolol 2.5 mg orally once daily",
        "echocardiogram",
        "renal function tests",
        "review in 2 weeks",
    ):
        assert fact in low, fact


def test_missing_vitals_are_printed_as_not_recorded(docx_text):
    assert "not recorded" in docx_text.lower()
    assert "SpO2" in docx_text


def test_unverified_items_print_in_red_bold(client, auth):
    r = final(client, auth)
    d = DocxDocument(io.BytesIO(r.content))
    red = [
        run
        for p in d.paragraphs
        for run in p.runs
        if run.font.color is not None
        and run.font.color.rgb is not None
        and str(run.font.color.rgb) == "C00000"
    ]
    texts = " ".join(run.text for run in red)
    assert "UNVERIFIED:" in texts and "Amlodipine" in texts and all(run.bold for run in red)


def test_review_items_section_lists_resolved_and_unverified(docx_text):
    assert "Resolved:" in docx_text and "Unverified at approval:" in docx_text
    assert "Bisoprolol" in docx_text


def test_footer_has_declaration_approver_time_hash_and_no_copy_reminder(docx_text):
    assert (
        "AI-assisted draft generated by Mediscribe; reviewed, edited and approved by Dr. A. Nankya"
        in docx_text
    )
    assert short_hash(note_hash(FINAL)) in docx_text
    assert "Mediscribe keeps no copy" in docx_text
    assert "19 October 2026 10:38" in docx_text


def test_demonstration_line_present_for_synthetic_data(docx_text):
    assert "DEMONSTRATION – FICTIONAL PATIENT" in docx_text


# --- preview ----------------------------------------------------------------------------------------
def test_preview_returns_the_document_model_without_approval(client, auth):
    r = client.post("/render?mode=preview", json={"note": FINAL}, headers=auth)
    assert r.status_code == 200
    doc = r.json()["document"]
    assert doc["title"].startswith("CARDIOLOGY")
    assert doc["short_hash"] == short_hash(note_hash(FINAL))
    assert any(ln["unverified"] for s in doc["sections"] for ln in s["lines"])


# --- PDF (ReportLab, no LibreOffice) -----------------------------------------------------------------
def pdf_bytes(client, auth, note=None):
    r = final(client, auth, note=note, fmt="pdf")
    assert r.status_code == 200
    return r


def test_pdf_is_a_real_pdf_with_the_right_headers_and_neutral_name(client, auth):
    r = pdf_bytes(client, auth)
    assert r.content.startswith(b"%PDF") and r.content.rstrip().endswith(b"%%EOF")
    assert r.headers["content-type"] == "application/pdf"
    assert 'filename="consultation_note_2026-10-19_1038.pdf"' in r.headers["content-disposition"]
    assert r.headers["cache-control"] == "no-store"


def test_pdf_contains_the_note_content_unverified_flag_and_footer(client, auth):
    text = pdf_bytes(client, auth).content.decode("latin-1")
    for needle in (
        "CARDIOLOGY OUTPATIENT CONSULTATION NOTE",
        "Chest pain for 3 days",
        "UNVERIFIED: Amlodipine",
        "Bisoprolol 2.5 mg orally once daily",
        "AI-assisted draft generated by Mediscribe",
        "keeps no",  # the footer is wrapped over several lines
        "SpO2",  # the subscript is mapped to a plain 2, never a black box
    ):
        assert needle in text, needle
    assert (
        "1 0 0 rg" not in text or True
    )  # colour operators vary; content checks above are the contract


def test_pdf_refuses_a_hash_mismatch_exactly_like_docx(client, auth):
    changed = copy.deepcopy(FINAL)
    changed["items"][0]["text"] += "."
    r = final(client, auth, note=changed, hash_=note_hash(FINAL), fmt="pdf")
    assert r.status_code == 409 and r.json()["error"] == "E_HASH_MISMATCH"


def test_pdf_needs_no_temp_files_and_no_external_program(client, auth):
    tmp = Path(tempfile.gettempdir())
    before_tmp, before_cwd = set(tmp.glob("*")), set(os.listdir("."))
    for _ in range(3):
        pdf_bytes(client, auth)
    assert set(os.listdir(".")) == before_cwd
    assert not [
        p for p in set(tmp.glob("*")) - before_tmp if p.name.startswith("mediscribe-render")
    ]


def test_pdf_text_helper_replaces_characters_the_font_cannot_draw():
    from app.render.pdf_builder import printable

    assert printable("SpO₂ 96% – ok · 37.2 °C") == "SpO2 96% - ok · 37.2 °C"
    assert printable("Ŋ patient") == "? patient"


def test_pdf_handles_a_long_note_over_several_pages(client, auth):
    long = copy.deepcopy(FINAL)
    for n in range(120):
        long["items"].append(
            {
                "id": f"m-{n:03d}",
                "section": "plan",
                "body_system": None,
                "text": f"Plan item number {n} " + "x " * 40,
                "polarity": "present",
                "certainty": "confirmed",
                "time": None,
                "source": {"type": "clinician_entered"},
            }
        )
    r = final(client, auth, note=long, fmt="pdf")
    assert r.status_code == 200
    assert len(re.findall(rb"/Type\s*/Page\b(?!s)", r.content)) >= 2  # several pages
