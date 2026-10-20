"""DOCX rendering with python-docx (Blueprint §4.3, §12.1). Built in memory; nothing touches disk."""

import io

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.render.model import Document

NAVY = RGBColor(0x1F, 0x4E, 0x79)
RED = RGBColor(0xC0, 0x00, 0x00)
GREY = RGBColor(0x59, 0x63, 0x73)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _run(par, text: str, *, bold=False, size=10, color=None, italic=False):
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = "Calibri"
    if color is not None:
        r.font.color.rgb = color
    return r


def render_docx(doc: Document) -> bytes:
    d = DocxDocument()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4 (§2.1)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.6)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    _run(p, doc.title, bold=True, size=15, color=NAVY)
    if doc.demo_line:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _run(p, doc.demo_line, bold=True, size=9, color=RED)

    table = d.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in doc.header_rows:
        row = table.add_row()
        row.cells[0].width, row.cells[1].width = Cm(3.2), Cm(14.2)
        _shade(row.cells[0], "E4EDF7")
        _run(row.cells[0].paragraphs[0], label, bold=True, size=9, color=NAVY)
        _run(row.cells[1].paragraphs[0], value, size=9.5)

    def heading(text: str):
        h = d.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(1)
        h.paragraph_format.keep_with_next = True
        _run(h, text, bold=True, size=10, color=NAVY)

    for s in doc.sections:
        heading(s.title)
        for ln in s.lines:
            par = d.add_paragraph()
            par.paragraph_format.space_after = Pt(1)
            if ln.unverified:
                _run(par, "UNVERIFIED: ", bold=True, color=RED)
                _run(par, ln.text, bold=True, color=RED)
            else:
                _run(par, ln.text)

    heading("REVIEW ITEMS")
    for ln in doc.review:
        par = d.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        _run(par, ln.text, bold=ln.unverified, color=RED if ln.unverified else None)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _run(footer, doc.footer, size=8, color=GREY, italic=True)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
